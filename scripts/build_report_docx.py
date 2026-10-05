"""Build an editable Word (.docx) version of the design report.

    python scripts/build_report_docx.py --name "Your Name" --id "Your ID"

Uses the same source as the PDF (docs/report/report_body.html + the generated appendix), so the
two versions always match. Output: docs/PawsConnect_Design_Report.docx
Needs:  pip install python-docx beautifulsoup4 pillow
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_report  # noqa: E402

REPORT_DIR = ROOT / "docs" / "report"
BROWN = RGBColor(0x7A, 0x3B, 0x17)


# ---------------------------------------------------------------- low-level helpers
def shade(element, fill: str) -> None:
    """Apply a background colour to a table cell or paragraph."""
    props = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def add_runs(par, node, bold=False, italic=False, mono=False, size=None) -> None:
    """Copy inline HTML (b, i, code, br, span.blank) into a docx paragraph."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = re.sub(r"\s+", " ", str(child))
            if not text:
                continue
            run = par.add_run(text)
            run.bold, run.italic = bold, italic
            if mono:
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            if size:
                run.font.size = Pt(size)
        elif isinstance(child, Tag):
            if child.name == "br":
                par.add_run().add_break()
            elif child.name == "span" and "blank" in (child.get("class") or []):
                par.add_run("_" * 22)
            else:
                add_runs(par, child,
                         bold=bold or child.name in ("b", "strong"),
                         italic=italic or child.name in ("i", "em"),
                         mono=mono or child.name == "code", size=size)


def add_table(doc, tag: Tag) -> None:
    rows = tag.find_all("tr")
    ncols = max(len(r.find_all(["th", "td"])) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for ri, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        for ci, cell in enumerate(cells):
            dc = table.cell(ri, ci)
            dc.text = ""
            par = dc.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            is_head = cell.name == "th"
            add_runs(par, cell, bold=is_head, size=8.5)
            if is_head:
                shade(dc, "F1E7DA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc, src: str) -> None:
    path = (REPORT_DIR / src).resolve()
    with Image.open(path) as im:
        aspect = im.height / im.width
    width = min(6.0, 8.1 / aspect)                   # keep tall screenshots on one page
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True


def add_page_number_footer(section) -> None:
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("PawsConnect design report - page ")
    run.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(8)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


# ---------------------------------------------------------------- walk the HTML
def convert(container: Tag, doc) -> None:
    for el in container.children:
        if not isinstance(el, Tag):
            continue
        name, classes = el.name, el.get("class") or []
        if name == "h1":
            p = doc.add_paragraph(style="Title")
            add_runs(p, el)
        elif name == "h2":
            p = doc.add_heading(level=1)
            add_runs(p, el)
        elif name == "h3":
            p = doc.add_heading(level=2)
            add_runs(p, el)
            if "pb" in classes:
                p.paragraph_format.page_break_before = True
        elif name == "p":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            add_runs(p, el)
        elif name == "table":
            add_table(doc, el)
        elif name == "figure":
            img = el.find("img")
            if img is not None:
                add_picture(doc, img["src"])
            cap = el.find("figcaption")
            if cap is not None:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(10)
                add_runs(p, cap, italic=False, size=8.5)
        elif name == "div" and "sub" in classes:
            p = doc.add_paragraph()
            add_runs(p, el, italic=True)
        elif name == "div" and "meta" in classes:
            p = doc.add_paragraph()
            add_runs(p, el)
            shade(p, "F6F1EA")
        elif name == "div" and "note" in classes:
            p = doc.add_paragraph()
            add_runs(p, el, size=9)
            shade(p, "FBF7F1")
        elif name == "div":                           # e.g. div.appendix wrapper
            convert(el, doc)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="")
    ap.add_argument("--id", default="")
    args = ap.parse_args()

    html_doc, _ = build_report.build_html(args.name, args.id)
    soup = BeautifulSoup(html_doc, "html.parser")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for style, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        st = doc.styles[style]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, BROWN
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title = doc.styles["Title"]
    title.font.name, title.font.size, title.font.color.rgb = "Calibri", Pt(21), RGBColor(0x3B, 0x2A, 0x1E)

    add_page_number_footer(sec)
    convert(soup.body, doc)

    out = ROOT / "docs" / "PawsConnect_Design_Report.docx"
    doc.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    main()
